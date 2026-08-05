# assignment.py

# import json
# json_dict = '{"name" : "Alice", "age" : 25, "city" : "Seoul"}'
# data = json.loads(json_dict)
# print(data)

# python_dict = {"name" : "Alice", "age" : 25, "city" : "Seoul"}
# json_output = json.dumps(python_dict, indent = 4)
# print(json_output)

# with open(r"ch21\data.json", "w", encoding = "utf-8") as f:
#     json.dump(python_dict, f, indent = 4)

# with open(r"ch21\data.json", "r", encoding = "utf-8") as f:
#     loaded_data = json.load(f)
# print(loaded_data)
# print(type(loaded_data))

# import os

# # 폴더 지정
# folder_path = r"C:\ROKEY\py_work\ch21"

# if os.path.exists(folder_path):
#     items = os.listdir(folder_path)
#     print(items)
# else:
#     print("해당 경로는 존재하지 않습니다.")

# from pathlib import Path
# path = r"ch21"
# folder_path = Path(path)
# txt_files = list(folder_path.glob("*.py"))
# for txt in txt_files:
#     print(txt.name)

# from pathlib import Path
# path = r"ch21"
# folder_path = Path(path)
# txt_files  = list(folder_path.glob("*.txt"))
# for txt in txt_files:
#     print(txt.name)

# from pathlib import Path
# import os

# new_folder = r"ch21\sample"
# if not os.path.exists(new_folder):
#     Path.mkdir(new_folder)
#     print("폴더 생성")

# from pathlib import Path
# import os
# new_folder = r"ch21\sample"
# if not os.path.exists(new_folder):
#     Path.mkdir(new_folder)
# else:
#     print(f"이미 {new_folder}에 폴더가 존재합니다.")

# from openpyxl import Workbook

# wb = Workbook()
# sheet = wb.active
# sheet.title = "Sheet1"

# path = r"ch21\data.xlsx"
# wb.save(path)

# from openpyxl import Workbook
# wb = Workbook()
# sheet = wb.active
# sheet.title = "Sheet1"
# save_path = r"ch21\data.xlsx"
# wb.save(save_path)


from openpyxl import load_workbook
path = r"ch21\data.xlsx"
wb = load_workbook(path)
sheet = wb.active
sheet.append(["이름" "나이" "지역"])
sheet.append(["김로키", 41, "부산"])
sheet.append(["최두산", 11, "서울"])
sheet.append(["이캠프", 24, "광주"])
wb.save(path)
